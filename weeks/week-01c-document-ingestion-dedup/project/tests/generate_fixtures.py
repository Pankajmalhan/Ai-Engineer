"""Generates real (non-fake) .docx and digital-text .pdf fixtures on disk via
python-docx and reportlab, so loader tests exercise unstructured's actual
docx/pdf partitioners against real files rather than hand-crafted mocks.
sample.md and sample.html are checked in directly (trivial to write by hand);
these two need a real library to produce a valid binary format.
"""

from __future__ import annotations

from pathlib import Path

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"

DOCX_PATH = FIXTURES_DIR / "sample.docx"
PDF_PATH = FIXTURES_DIR / "sample_digital.pdf"

DOCX_HEADING = "Incident Response Runbook"
DOCX_BODY = (
    "Page the on-call engineer within five minutes of a triggered alert. "
    "File a written postmortem within three business days of resolution."
)
DOCX_CONTACT = "Escalate to oncall-lead@example.com if unacknowledged."

PDF_TITLE = "Quarterly Infrastructure Report"
PDF_BODY = "Compute spend increased twelve percent due to nightly model retraining."


def ensure_fixtures() -> None:
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    if not DOCX_PATH.exists():
        _write_docx()
    if not PDF_PATH.exists():
        _write_pdf()


def _write_docx() -> None:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading(DOCX_HEADING, level=1)
    doc.add_paragraph(DOCX_BODY)
    doc.add_paragraph(DOCX_CONTACT)
    doc.save(DOCX_PATH)


def _write_pdf() -> None:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(PDF_PATH))
    c.drawString(72, 750, PDF_TITLE)
    c.drawString(72, 720, PDF_BODY)
    c.save()


if __name__ == "__main__":
    ensure_fixtures()
    print(f"Wrote {DOCX_PATH}")
    print(f"Wrote {PDF_PATH}")
